"""本地文件文本抽取：PDF、PPTX、DOCX 与纯文本，供 AI 阅读课件内容。"""

from __future__ import annotations

from pathlib import Path

from .textutil import clip, html_to_text

MAX_CHARS = 60_000
_TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".py", ".m", ".r", ".tex", ".log"}


def parse_pages(spec: str | None, total: int) -> list[int]:
    """把「1-5,8」解析为 0 起始的页码列表；空值表示全部。"""
    if not spec:
        return list(range(total))
    result: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, _, end = part.partition("-")
            lo = int(start) if start.strip() else 1
            hi = int(end) if end.strip() else total
            result.update(range(max(lo, 1) - 1, min(hi, total)))
        else:
            n = int(part)
            if 1 <= n <= total:
                result.add(n - 1)
    return sorted(result)


def _pdf(path: Path, pages: str | None) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    total = len(reader.pages)
    out = [f"PDF，共 {total} 页"]
    for i in parse_pages(pages, total):
        text = (reader.pages[i].extract_text() or "").strip()
        out.append(f"\n## 第 {i + 1} 页\n{text or '（本页无可抽取文字，可能是图片）'}")
    return "\n".join(out)


def _shape_text(shape) -> list[str]:  # noqa: ANN001 - python-pptx 动态类型
    lines: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if text:
                lines.append(("  " * para.level) + text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
    for child in getattr(shape, "shapes", []) or []:
        lines.extend(_shape_text(child))
    return lines


def _pptx(path: Path, pages: str | None) -> str:
    from pptx import Presentation

    slides = list(Presentation(str(path)).slides)
    out = [f"PPT，共 {len(slides)} 页"]
    for i in parse_pages(pages, len(slides)):
        slide = slides[i]
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[讲者备注] {notes}")
        out.append(f"\n## 第 {i + 1} 页\n" + ("\n".join(lines) or "（本页无文字，可能是图片）"))
    return "\n".join(out)


def _docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
    return "\n".join(lines)


def extract_text(path: Path, pages: str | None = None, max_chars: int = MAX_CHARS) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        body = _pdf(path, pages)
    elif ext in (".pptx", ".pptm", ".ppsx"):
        body = _pptx(path, pages)
    elif ext in (".docx", ".docm"):
        body = _docx(path)
    elif ext in (".html", ".htm"):
        body = html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    elif ext in _TEXT_EXTS:
        body = path.read_text(encoding="utf-8", errors="replace")
    elif ext in (".ppt", ".doc"):
        return f"{path.name} 是旧版 Office 格式，无法直接抽取文字；可用 Keynote/PowerPoint 另存为 .pptx/.pdf。"
    else:
        return f"{path.name} 不是可抽取文字的格式，已保存在：{path}"
    hint = "；内容过长时可用 pages 参数分段读取" if ext in (".pdf", ".pptx", ".pptm", ".ppsx") else ""
    return clip(f"# {path.name}\n本地路径：{path}{hint}\n\n{body}", max_chars)
